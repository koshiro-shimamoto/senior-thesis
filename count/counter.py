import argparse
from pathlib import Path
from collections import Counter
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

SWC_PATTERN = re.compile(r"SWC ID:\s*\d+")


def collect_swc_from_file(path: Path) -> Counter:
    counter = Counter()
    with path.open(encoding="utf-8", errors="ignore") as f:
        for line in f:
            match = SWC_PATTERN.search(line)
            if match:
                counter[match.group()] += 1
    return counter


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("before_dir", type=Path)
    parser.add_argument("after_dir", type=Path)
    parser.add_argument("--output", default="swc_comparison.xlsx")
    args = parser.parse_args()

    wb = Workbook()
    ws = wb.active
    ws.title = "SWC Comparison"

    current_row = 1

    for before_file in sorted(args.before_dir.iterdir()):
        if "error" in before_file.name:
            continue
        if not before_file.is_file() or not before_file.name.endswith(".txt"):
            continue

        after_file = args.after_dir / before_file.name.replace(".txt", ".bin.txt")
        if not after_file.exists():
            continue

        before_counts = collect_swc_from_file(before_file)
        after_counts  = collect_swc_from_file(after_file)
        all_swcs = sorted(set(before_counts) | set(after_counts))

        if not all_swcs:
            continue

        # ===== ファイル名（セル結合） =====
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=4)
        cell = ws.cell(row=current_row, column=1)
        cell.value = before_file.name
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

        current_row += 1

        # ===== ヘッダ行 =====
        ws.append(["SWC ID", "Before", "After", "Diff"])
        for col in range(1, 5):
            ws.cell(row=current_row, column=col).font = Font(bold=True)

        current_row += 1

        # ===== データ行 =====
        for swc in all_swcs:
            b = before_counts.get(swc, 0)
            a = after_counts.get(swc, 0)
            diff = a - b
            ws.append([int(swc.replace("SWC ID:", "").strip()), b, a, diff])
            current_row += 1

        # ===== 空行（区切り） =====
        current_row += 1

    wb.save(args.output)
    print(f"[INFO] Excel file saved: {args.output}")


if __name__ == "__main__":
    main()
